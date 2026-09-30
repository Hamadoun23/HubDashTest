"""Classeur des comptes GDA Hub (données locales réelles), un onglet par application et type d'utilisateur."""
import json
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys
D = (sys.argv[1] if len(sys.argv) > 1 else "comptes").rstrip("/") + "/"
SORTIE = sys.argv[2] if len(sys.argv) > 2 else "DonneeEnProd/Comptes_GDA_Hub.xlsx"
MDP = sys.argv[3] if len(sys.argv) > 3 else "1234"

rapport = json.load(open(D + "rapport.json", encoding="utf-8"))
admins = json.load(open(D + "admins.json", encoding="utf-8"))

ORANGE, BRUN, CREME = "C8521A", "381419", "FBF1E4"
entete_font = Font(bold=True, color="FFFFFF")
entete_fill = PatternFill("solid", fgColor=BRUN)
titre_font = Font(bold=True, size=14, color=BRUN)
groupe_fill = PatternFill("solid", fgColor="F6E3CF")
fin = Side(style="thin", color="E4D6C6")
bord = Border(bottom=fin)

LIBELLES = {
    "rh": {"DIRECTION": "Direction", "RH": "Ressources humaines", "FINANCE": "Finance", "SALARIE": "Salariés"},
    "orange": {"admin": "Administrateurs", "direction": "Direction", "responsable_production": "Responsables production", "commercial": "Commerciaux", "finance": "Finance"},
    "daily": {"admin": "Administrateurs", "partenaire": "Partenaires (client)", "chef_chantier": "Chefs de chantier", "ingenieur": "Ingénieurs"},
    "planning": {"admin": "Administrateurs", "team": "Équipe", "client": "Clients"},
    "campagnes": {"admin": "Administrateurs", "direction": "Direction", "commercial": "Commerciaux", "commercial_telephonique": "Commerciaux téléphoniques"},
}
ORDRE = {
    "rh": ["DIRECTION", "RH", "FINANCE", "SALARIE"],
    "orange": ["admin", "direction", "responsable_production", "finance", "commercial"],
    "daily": ["admin", "chef_chantier", "ingenieur", "partenaire"],
    "planning": ["admin", "team", "client"],
    "campagnes": ["admin", "direction", "commercial", "commercial_telephonique"],
}

wb = Workbook()


def onglet(nom, titre, sous_titre, colonnes, groupes):
    """groupes : liste de (libellé du groupe, lignes)."""
    ws = wb.create_sheet(nom[:31])
    ws.sheet_view.showGridLines = False
    ws["A1"] = titre
    ws["A1"].font = titre_font
    ws["A2"] = sous_titre
    ws["A2"].font = Font(italic=True, color="7A6A5A")
    ligne = 4
    for c, libelle in enumerate(colonnes, 1):
        cell = ws.cell(row=ligne, column=c, value=libelle)
        cell.font, cell.fill = entete_font, entete_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[ligne].height = 20
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)
    for libelle, lignes in groupes:
        if not lignes:
            continue
        ligne += 1
        cell = ws.cell(row=ligne, column=1, value=f"{libelle} ({len(lignes)})")
        cell.font = Font(bold=True, color=BRUN)
        for c in range(1, len(colonnes) + 1):
            ws.cell(row=ligne, column=c).fill = groupe_fill
        for valeurs in lignes:
            ligne += 1
            for c, v in enumerate(valeurs, 1):
                cell = ws.cell(row=ligne, column=c, value=v)
                cell.border = bord
                if c == 1:
                    cell.number_format = "@"
    largeurs = [34, 12, 34, 26, 44]
    for c in range(1, len(colonnes) + 1):
        ws.column_dimensions[get_column_letter(c)].width = largeurs[c - 1] if c <= len(largeurs) else 20
    return ws


COLS = ["Identifiant de connexion", "Mot de passe", "Nom", "Rôle", "Détails"]
par_app = defaultdict(list)
for app, ident, nom, role, det in rapport:
    par_app[app].append((ident, MDP, nom, LIBELLES.get(app, {}).get(role, role), det, role))


def groupes_de(app, filtre=lambda l: True):
    lignes = [l for l in par_app[app] if filtre(l)]
    sortie = []
    for role in ORDRE[app]:
        sel = sorted([l[:5] for l in lignes if l[5] == role], key=lambda x: (x[2] or "").lower())
        sortie.append((LIBELLES[app].get(role, role), sel))
    autres = sorted([l[:5] for l in lignes if l[5] not in ORDRE[app]], key=lambda x: (x[2] or "").lower())
    sortie.append(("Autres", autres))
    return sortie


# --- Récapitulatif -----------------------------------------------------------
ws = wb.active
ws.title = "Récapitulatif"
ws.sheet_view.showGridLines = False
ws["A1"] = "Comptes GDA Hub — environnement local (données réelles)"
ws["A1"].font = titre_font
ws["A2"] = f"Mot de passe unique pour tous les comptes : {MDP}   ·   Connexion : http://localhost:8080"
ws["A2"].font = Font(bold=True, color=ORANGE)
ws["A3"] = "Données personnelles réelles : ne pas diffuser. Les commerciaux Campagnes se connectent avec leur numéro de téléphone."
ws["A3"].font = Font(italic=True, color="7A6A5A")
ws.cell(row=5, column=1, value="Application").font = entete_font
ws.cell(row=5, column=2, value="Type d'utilisateur").font = entete_font
ws.cell(row=5, column=3, value="Nombre").font = entete_font
for c in range(1, 4):
    ws.cell(row=5, column=c).fill = entete_fill
NOMS = {"rh": "RH & Finance", "orange": "Jus d'orange", "daily": "Chantiers", "planning": "Planning", "campagnes": "Campagnes"}
r = 6
for app in ["rh", "orange", "daily", "planning", "campagnes"]:
    for libelle, lignes in groupes_de(app):
        if lignes:
            ws.cell(row=r, column=1, value=NOMS[app])
            ws.cell(row=r, column=2, value=libelle)
            ws.cell(row=r, column=3, value=len(lignes))
            r += 1
ws.cell(row=r, column=1, value="Administrateurs du hub").font = Font(bold=True)
ws.cell(row=r, column=3, value=len(admins))
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 10

# --- Onglets ------------------------------------------------------------------
onglet("Administrateurs hub", "Administrateurs du hub", "Accès à l'administration de GDA Hub (comptes, droits, applications).",
       COLS, [("Administrateurs", [(i, MDP, n, "Super-administrateur" if s else "Administrateur", ", ".join(h)) for i, n, s, h in admins])])
onglet("RH & Finance", "RH & Finance", "Agents de FinanceRH (congés, présences, demandes, finance).", COLS, groupes_de("rh"))
onglet("Jus d'orange", "Jus d'orange", "Utilisateurs de l'application Jus d'orange.", COLS, groupes_de("orange"))
onglet("Chantiers", "Chantiers (daily)", "Suivi de chantier — le partenaire voit la vue client.", COLS, groupes_de("daily"))
onglet("Planning", "GDA Media Planning", "Équipe média et comptes clients.", COLS, groupes_de("planning"))
est_bdm = lambda l: l[4].startswith("BDM")
est_uba = lambda l: l[4].startswith("UBA")
interne = lambda l: l[5] in ("admin", "direction")
onglet("Campagnes - Admin", "Campagnes — administration et direction", "Choisissent le client (BDM, UBA) à la connexion.",
       COLS, [(g, l) for g, l in groupes_de("campagnes", interne)])
onglet("Campagnes - BDM", "Campagnes — commerciaux BDM", "Commerciaux externes : identifiant = numéro de téléphone. Ne voient que Campagnes.",
       COLS, [(g, l) for g, l in groupes_de("campagnes", lambda l: est_bdm(l) and not interne(l))])
onglet("Campagnes - UBA", "Campagnes — commerciaux UBA", "Commerciaux externes : identifiant = numéro de téléphone. Ne voient que Campagnes.",
       COLS, [(g, l) for g, l in groupes_de("campagnes", lambda l: est_uba(l) and not interne(l))])

# Tous les comptes : une ligne par compte, ses applications.
comptes = defaultdict(lambda: {"nom": "", "apps": []})
for app, ident, nom, role, det in rapport:
    comptes[ident]["nom"] = comptes[ident]["nom"] or nom
    comptes[ident]["apps"].append(f"{NOMS[app]} ({LIBELLES.get(app, {}).get(role, role)})")
onglet("Tous les comptes", "Tous les comptes", "Une ligne par compte du hub, avec ses applications.",
       ["Identifiant de connexion", "Mot de passe", "Nom", "Nb d'applications", "Applications"],
       [("Comptes", sorted([(i, MDP, c["nom"], len(c["apps"]), " · ".join(c["apps"])) for i, c in comptes.items()], key=lambda x: (x[2] or "").lower()))])

wb.save(SORTIE)
print("Classeur écrit :", SORTIE, "| onglets :", wb.sheetnames, "| comptes :", len(comptes))
